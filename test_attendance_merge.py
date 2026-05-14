import unittest
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import load_workbook

from attendance_parser import generate_report, merge_cross_source


def record(emp_id, name, source, am_start='', am_end='', pm_start='', pm_end=''):
    return {
        '员工号': emp_id,
        '姓名': name,
        '部门': '仓库',
        '日期': '4月20日',
        '星期': '一',
        '上午上班': am_start,
        '上午下班': am_end,
        '下午上班': pm_start,
        '下午下班': pm_end,
        '加班签到': '',
        '加班签退': '',
        '数据来源': source,
    }


class AttendanceMergeTest(unittest.TestCase):
    def test_does_not_merge_employee_id_collision_with_different_names(self):
        df = pd.DataFrame([
            record(23, 'JERRY', 'warehouse-a', am_start='08:50:00', am_end='12:00:00'),
            record(23, 'OTHER PERSON', 'warehouse-b', pm_start='13:00:00', pm_end='18:00:00'),
        ])

        merged = merge_cross_source(df)

        self.assertEqual(len(merged), 2)
        jerry = merged[merged['姓名'] == 'JERRY'].iloc[0]
        self.assertEqual(jerry['下班打卡'], '12:00:00')
        self.assertEqual(jerry['数据来源'], 'warehouse-a')

    def test_keeps_existing_same_name_merge_when_employee_ids_differ(self):
        df = pd.DataFrame([
            record(2001, 'ISABEL', 'warehouse-a', am_start='09:00:00', am_end='12:00:00'),
            record(3001, 'ISABEL', 'warehouse-b', pm_start='13:00:00', pm_end='17:00:00'),
        ])

        merged = merge_cross_source(df)

        self.assertEqual(len(merged), 1)
        self.assertEqual(merged.iloc[0]['合计工时'], 8)

    def test_merges_different_names_with_same_mapping(self):
        df = pd.DataFrame([
            record(2001, 'YING YU', 'warehouse-a', am_start='09:00:00', am_end='12:00:00'),
            record(3001, 'YINGYU', 'warehouse-b', pm_start='13:00:00', pm_end='17:00:00'),
        ])
        mapping = {
            'YING YU': {'first_name': 'YING', 'last_name': 'YU'},
            'YINGYU': {'first_name': 'YING', 'last_name': 'YU'},
        }

        merged = merge_cross_source(df, name_mapping=mapping)

        self.assertEqual(len(merged), 1)
        self.assertEqual(merged.iloc[0]['姓名'], 'YING YU')
        self.assertEqual(merged.iloc[0]['合计工时'], 8)

    def test_adjustment_table_matches_any_merged_alias(self):
        df = pd.DataFrame([
            record(1001, 'YING YU', 'warehouse-a', am_start='08:00:00', am_end='12:00:00'),
            record(1001, 'YINGYU', 'warehouse-b', pm_start='13:00:00', pm_end='18:00:00'),
        ])
        mapping = {
            'YING YU': {'first_name': 'YING', 'last_name': 'YU'},
            'YINGYU': {'first_name': 'YING', 'last_name': 'YU'},
        }

        merged = merge_cross_source(
            df,
            adj_map={('YINGYU', '04月20日'): 7},
            name_mapping=mapping,
        )

        self.assertEqual(len(merged), 1)
        self.assertEqual(merged.iloc[0]['上班打卡'], '07:00:00')
        self.assertEqual(merged.iloc[0]['合计工时'], 11)
        self.assertEqual(merged.iloc[0]['调整标记'], '已调整')

    def test_weekly_report_counts_only_positive_hour_days(self):
        rows = [
            {
                '员工号': 23, '姓名': 'JERRY', '部门': '仓库', '日期': '4月20日', '星期': '一',
                '上班打卡': '09:00:00', '下班打卡': '18:00:00',
                '上午上班': '09:00:00', '上午下班': '', '下午上班': '', '下午下班': '',
                '加班签到': '', '加班签退': '',
                '工作工时': 9, '加班工时': 0, '合计工时': 9, '调整标记': '', '数据来源': 'warehouse-a',
            },
            {
                '员工号': 23, '姓名': 'JERRY', '部门': '仓库', '日期': '4月21日', '星期': '二',
                '上班打卡': '', '下班打卡': '',
                '上午上班': '', '上午下班': '', '下午上班': '', '下午下班': '',
                '加班签到': '', '加班签退': '',
                '工作工时': 0, '加班工时': 0, '合计工时': 0, '调整标记': '', '数据来源': 'warehouse-a',
            },
        ]

        with TemporaryDirectory() as tmpdir:
            output_path = f'{tmpdir}/report.xlsx'
            generate_report(pd.DataFrame(rows), output_path)
            ws = load_workbook(output_path, read_only=True, data_only=True)['每周工时汇总']
            header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            values = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]

        self.assertEqual(values[header.index('打卡天数')], 1)


if __name__ == '__main__':
    unittest.main()
